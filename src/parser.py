from docx import Document


def parse_resume(file_path):
    document = Document(file_path)

    paragraphs = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            style_name = None

            if paragraph.style is not None:
                style_name = paragraph.style.name

            paragraphs.append({
                "text": text,
                "style": style_name
            })

    return {
        "paragraphs": paragraphs,
        "tables": document.tables
    }


def get_resume_text(parsed_resume):
    return "\n".join(
        paragraph["text"]
        for paragraph in parsed_resume["paragraphs"]
    )
def load_job_posting(file_path):
  with open(file_path, "r", encoding="utf-8") as file:
    return file.read()
  