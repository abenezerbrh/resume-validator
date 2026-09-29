import json
from datetime import datetime
from pathlib import Path

import pytest
from docx import Document

from src.resume_versioning import (
    METADATA_FILENAME,
    VersionExistsError,
    get_version_dir,
    job_id_from_path,
    make_job_id,
    save_tailored_version
)


CREATED_AT = datetime(2026, 9, 29, 14, 30, 5)

TAILORED_TEXT = "Experience\nDeveloped an internal Python tool.\n"


def create_resume(path):
    document = Document()
    document.add_paragraph().add_run("Experience").bold = True
    document.add_paragraph("Built an internal tool in Python.")
    document.save(path)


def test_creates_tailored_resume_version_with_metadata(tmp_path):
    original = tmp_path / "Abenezer_Balcha_Resume_BMO.docx"
    create_resume(original)

    version = save_tailored_version(
        original,
        TAILORED_TEXT,
        "BMO",
        created_at=CREATED_AT,
        output_dir=tmp_path / "versions"
    )

    docx_path = Path(version["tailored_docx"])
    assert docx_path.exists()
    assert [p.text for p in Document(docx_path).paragraphs] == [
        "Experience",
        "Developed an internal Python tool."
    ]

    assert version["job_id"] == "BMO"
    assert version["version_id"] == "20260929-143005"
    assert version["created_at"] == "2026-09-29T14:30:05"
    assert version["original_resume"] == "Abenezer_Balcha_Resume_BMO.docx"

    metadata_path = docx_path.parent / METADATA_FILENAME
    assert json.loads(metadata_path.read_text(encoding="utf-8")) == version


def test_generates_expected_versioned_output_path(tmp_path):
    original = tmp_path / "Abenezer_Balcha_Resume_BMO.docx"
    create_resume(original)
    versions = tmp_path / "versions"

    version = save_tailored_version(
        original,
        TAILORED_TEXT,
        job_id_from_path("jobs/software_developer.txt"),
        created_at=CREATED_AT,
        output_dir=versions
    )

    assert Path(version["tailored_docx"]) == (
        versions / "software_developer" / "20260929-143005"
        / "Abenezer_Balcha_Resume_BMO_tailored.docx"
    )


def test_job_id_is_folder_safe():
    assert job_id_from_path("jobs/software_developer.txt") == "software_developer"
    assert make_job_id("BMO") == "BMO"
    assert make_job_id("RBC / Dev Role") == "RBC_Dev_Role"


def test_does_not_overwrite_existing_version(tmp_path):
    original = tmp_path / "resume.docx"
    create_resume(original)
    versions = tmp_path / "versions"

    first = save_tailored_version(
        original, TAILORED_TEXT, "BMO",
        created_at=CREATED_AT, output_dir=versions
    )
    first_docx = Path(first["tailored_docx"])
    first_metadata = first_docx.parent / METADATA_FILENAME
    docx_bytes = first_docx.read_bytes()
    metadata_bytes = first_metadata.read_bytes()

    with pytest.raises(VersionExistsError):
        save_tailored_version(
            original, "Education\nSomething different\n", "BMO",
            created_at=CREATED_AT, output_dir=versions
        )

    assert first_docx.read_bytes() == docx_bytes
    assert first_metadata.read_bytes() == metadata_bytes


def test_new_version_is_saved_alongside_existing_one(tmp_path):
    original = tmp_path / "resume.docx"
    create_resume(original)
    versions = tmp_path / "versions"

    first = save_tailored_version(
        original, TAILORED_TEXT, "BMO",
        created_at=CREATED_AT, output_dir=versions
    )
    first_bytes = Path(first["tailored_docx"]).read_bytes()

    second = save_tailored_version(
        original, TAILORED_TEXT, "BMO",
        created_at=datetime(2026, 9, 29, 15, 0, 0), output_dir=versions
    )

    assert first["tailored_docx"] != second["tailored_docx"]
    assert Path(first["tailored_docx"]).read_bytes() == first_bytes
    assert Path(second["tailored_docx"]).exists()


def test_original_resume_is_unchanged(tmp_path):
    original = tmp_path / "resume.docx"
    create_resume(original)
    original_bytes = original.read_bytes()

    version = save_tailored_version(
        original, TAILORED_TEXT, "BMO",
        created_at=CREATED_AT, output_dir=tmp_path / "versions"
    )

    assert original.read_bytes() == original_bytes
    assert Path(version["tailored_docx"]) != original


def test_failed_generation_leaves_no_partial_version(tmp_path):
    original = tmp_path / "resume.docx"
    create_resume(original)
    versions = tmp_path / "versions"

    def failing_generate(*args):
        raise RuntimeError("generation failed")

    with pytest.raises(RuntimeError):
        save_tailored_version(
            original, TAILORED_TEXT, "BMO",
            created_at=CREATED_AT, output_dir=versions,
            generate=failing_generate
        )

    assert not get_version_dir("BMO", "20260929-143005", versions).exists()
