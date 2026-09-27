from pathlib import Path

import pytest
from docx import Document
from docx.shared import Inches, Pt

from src.docx_generator import (
    DocxGenerationError,
    generate_tailored_docx,
    get_tailored_output_path
)


def create_resume(path):
    document = Document()
    section = document.sections[0]
    section.left_margin = Inches(0.4)
    section.top_margin = Inches(0.3)

    heading = document.add_paragraph().add_run("Experience")
    heading.bold = True
    heading.font.size = Pt(12)

    job_line = document.add_paragraph()
    company = job_line.add_run("Acme Corp")
    company.bold = True
    location = job_line.add_run("\tToronto, ON")
    location.bold = False

    bullet = document.add_paragraph(style="List Bullet")
    bullet.add_run("Built an internal tool in Python.").font.size = Pt(10.5)

    document.save(path)


def test_generates_tailored_docx_with_original_formatting(tmp_path):
    original = tmp_path / "resume.docx"
    output = tmp_path / "out" / "resume_tailored.docx"
    create_resume(original)
    original_bytes = original.read_bytes()

    tailored_text = (
        "Experience\n"
        "Acme Corporation\tToronto, ON\n"
        "Developed an internal Python tool used by the team.\n"
    )

    generate_tailored_docx(original, tailored_text, output)

    assert original.read_bytes() == original_bytes

    document = Document(output)
    paragraphs = document.paragraphs

    assert [p.text for p in paragraphs] == [
        "Experience",
        "Acme Corporation\tToronto, ON",
        "Developed an internal Python tool used by the team."
    ]

    assert document.sections[0].left_margin == Inches(0.4)
    assert document.sections[0].top_margin == Inches(0.3)

    heading_run = paragraphs[0].runs[0]
    assert heading_run.bold is True
    assert heading_run.font.size == Pt(12)

    company_run, location_run = paragraphs[1].runs
    assert company_run.text == "Acme Corporation"
    assert company_run.bold is True
    assert location_run.bold is False

    assert paragraphs[2].style.name == "List Bullet"
    assert paragraphs[2].runs[0].font.size == Pt(10.5)


def test_bolds_technical_skills_labels_only(tmp_path):
    original = tmp_path / "resume.docx"
    output = tmp_path / "resume_tailored.docx"

    document = Document()
    for heading_text in ("Technical Skills", "Education"):
        heading = document.add_paragraph().add_run(heading_text)
        heading.bold = True
        heading.font.size = Pt(12)
        document.add_paragraph().add_run("Languages: Python, SQL").font.size = Pt(10.5)
    document.save(original)

    tailored_text = (
        "Technical Skills\n"
        "Languages: Python, SQL, JavaScript\n"
        "Education\n"
        "Languages: Python, SQL\n"
    )

    generate_tailored_docx(original, tailored_text, output)

    paragraphs = Document(output).paragraphs

    label_run, items_run = paragraphs[1].runs
    assert label_run.text == "Languages:"
    assert label_run.bold is True
    assert items_run.text == " Python, SQL, JavaScript"
    assert not items_run.bold
    assert items_run.font.size == Pt(10.5)

    # Lines outside the Technical Skills section are left alone.
    assert len(paragraphs[3].runs) == 1
    assert not paragraphs[3].runs[0].bold


def test_refuses_to_overwrite_original(tmp_path):
    original = tmp_path / "resume.docx"
    create_resume(original)

    with pytest.raises(DocxGenerationError):
        generate_tailored_docx(original, "Experience", original)


def test_output_path_is_marked_tailored():
    assert get_tailored_output_path(
        "resumes/Abenezer_Balcha_Resume_BMO.docx"
    ) == Path("reports/Abenezer_Balcha_Resume_BMO_tailored.docx")
